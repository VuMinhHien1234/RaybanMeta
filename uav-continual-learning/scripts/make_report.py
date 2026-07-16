#!/usr/bin/env python3
"""Sinh báo cáo kết quả từ artifacts/results/ -> artifacts/BAO_CAO_KET_QUA.docx

  python scripts/make_report.py            # đọc ./artifacts
  python scripts/make_report.py --dir path/khac

Nội dung: bảng toàn bộ run · biểu đồ Acc/Forgetting theo method từng dataset ·
kết luận tự động (method tốt nhất, HOPE có >1+1 không, khoảng cách tới replay/ncm).
Thiếu python-docx -> tự fallback ra BAO_CAO_KET_QUA.md.
"""
from __future__ import annotations

import argparse
import json
import pathlib
import platform
import sys
from datetime import datetime

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1] / "src"))

COLS = ["dataset", "method", "optimizer", "backbone", "seed",
        "average_accuracy", "average_forgetting", "backward_transfer",
        "forward_transfer", "trainable_params", "method_extra_floats", "runtime_sec"]


def load_rows(root: pathlib.Path) -> list[dict]:
    rows = []
    for p in sorted(root.glob("*/metrics.json")):
        m = json.loads(p.read_text(encoding="utf-8"))
        m["run"] = p.parent.name
        rows.append(m)
    return [r for r in rows if r.get("dataset") != "synthetic"]  # bỏ smoke


def label(r: dict) -> str:
    lab = r["method"]
    if r.get("optimizer", "adamw") != "adamw":
        lab += f"+{r['optimizer']}"
    if r.get("cms_order"):
        lab += f"({r['cms_order'][:5]},{r.get('cms_periods', '')})"
    if "titans" in r["method"] and "_never" not in r["run"]:
        # phân biệt các bậc reset của titans qua tên run nếu có
        for tag in ("image", "task", "never"):
            if r["run"].endswith(tag):
                lab += f"[{tag}]"
    return lab


def make_charts(rows, out_dir: pathlib.Path) -> list[pathlib.Path]:
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except ImportError:
        return []
    out_dir.mkdir(parents=True, exist_ok=True)
    paths = []
    for ds in sorted({r["dataset"] for r in rows}):
        sub = sorted([r for r in rows if r["dataset"] == ds], key=lambda r: -r["average_accuracy"])
        labs = [label(r) for r in sub]
        fig, axes = plt.subplots(1, 2, figsize=(13, 4.2))
        axes[0].barh(labs[::-1], [r["average_accuracy"] for r in sub][::-1], color="#2b8a3e")
        axes[0].set_title(f"{ds} — Average Accuracy (cao = tốt)")
        axes[0].set_xlim(0, 1)
        axes[1].barh(labs[::-1], [r["average_forgetting"] for r in sub][::-1], color="#c92a2a")
        axes[1].set_title(f"{ds} — Forgetting (thấp = tốt)")
        fig.tight_layout()
        p = out_dir / f"chart_{ds}.png"
        fig.savefig(p, dpi=140)
        plt.close(fig)
        paths.append(p)
    return paths


def auto_findings(rows) -> list[str]:
    finds = []
    for ds in sorted({r["dataset"] for r in rows}):
        sub = [r for r in rows if r["dataset"] == ds]
        best = max(sub, key=lambda r: r["average_accuracy"])
        finds.append(f"[{ds}] Tốt nhất về Accuracy: {label(best)} "
                     f"(Acc {best['average_accuracy']:.4f}, Forgetting {best['average_forgetting']:.4f}).")
        least = min(sub, key=lambda r: r["average_forgetting"])
        finds.append(f"[{ds}] Ít quên nhất: {label(least)} (Forgetting {least['average_forgetting']:.4f}).")

        def get(name):
            c = [r for r in sub if r["method"] == name]
            return max(c, key=lambda r: r["average_accuracy"]) if c else None

        hope, cms, tit = get("hope"), get("cms"), get("titans")
        if hope and cms and tit:
            better = hope["average_accuracy"] > max(cms["average_accuracy"], tit["average_accuracy"])
            finds.append(f"[{ds}] HOPE {'VƯỢT' if better else 'CHƯA vượt'} cả CMS-only lẫn Titans-only "
                         f"(HOPE {hope['average_accuracy']:.4f} vs CMS {cms['average_accuracy']:.4f} "
                         f"vs Titans {tit['average_accuracy']:.4f}) -> {'1+1>2' if better else 'cần phân tích thành phần'}.")
        replay, ncm = get("replay"), get("ncm")
        nl_best = max((r for r in sub if r["method"] in ("titans", "cms", "hope")),
                      key=lambda r: r["average_accuracy"], default=None)
        if nl_best and replay and ncm:
            finds.append(f"[{ds}] NL tốt nhất ({label(nl_best)}, {nl_best['average_accuracy']:.4f}) so với mốc: "
                         f"NCM {ncm['average_accuracy']:.4f} ({'ĐÃ vượt' if nl_best['average_accuracy'] > ncm['average_accuracy'] else 'chưa vượt'}), "
                         f"replay {replay['average_accuracy']:.4f} ({'ĐÃ vượt' if nl_best['average_accuracy'] > replay['average_accuracy'] else 'chưa vượt'}, "
                         f"nhưng NL không lưu ảnh thô: {nl_best.get('method_extra_floats', 0):,} vs {replay.get('method_extra_floats', 0):,} floats).")
    return finds


def fmt(v, key):
    if v is None:
        return "—"
    if key in ("average_accuracy", "average_forgetting", "backward_transfer", "forward_transfer"):
        return f"{float(v):.4f}"
    if key in ("trainable_params", "method_extra_floats"):
        return f"{int(v):,}"
    if key == "runtime_sec":
        return f"{float(v):.0f}"
    return str(v)


def write_docx(rows, charts, finds, out_path: pathlib.Path) -> bool:
    try:
        from docx import Document
        from docx.shared import Inches, Pt
    except ImportError:
        return False
    doc = Document()
    doc.add_heading("UAV Continual Learning — Báo cáo kết quả tự động", level=0)
    gpu = "n/a"
    try:
        import torch
        gpu = torch.cuda.get_device_name(0) if torch.cuda.is_available() else "CPU/MPS"
    except Exception:
        pass
    doc.add_paragraph(f"Sinh lúc: {datetime.now():%Y-%m-%d %H:%M} · máy: {platform.node()} · thiết bị: {gpu}. "
                      f"Mọi số tái lập từ artifacts/results/<run>/ (config.yaml + acc_matrix.csv đi kèm từng run).")

    doc.add_heading("1. Bảng tổng hợp toàn bộ run", level=1)
    table = doc.add_table(rows=1, cols=len(COLS))
    table.style = "Light Grid Accent 1"
    for i, c in enumerate(COLS):
        table.rows[0].cells[i].text = c
    for r in sorted(rows, key=lambda x: (x["dataset"], -x["average_accuracy"])):
        cells = table.add_row().cells
        for i, c in enumerate(COLS):
            cells[i].text = fmt(r.get(c), c)
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(8)

    doc.add_heading("2. Biểu đồ so sánh", level=1)
    for p in charts:
        doc.add_picture(str(p), width=Inches(6.5))

    doc.add_heading("3. Kết luận tự động (kiểm tra lại trước khi dùng trong báo cáo!)", level=1)
    for f in finds:
        doc.add_paragraph(f, style="List Bullet")

    doc.add_heading("4. Việc tiếp theo", level=1)
    doc.add_paragraph("Đối chiếu các cổng chuyển trong plans/ (G2 §6, G3 §6, TASKS_G4 §D); đọc log "
                      "norm(state) và ‖Δw‖ trong run.log; chạy đa seed ở G5 trước khi kết luận cuối.")
    doc.save(out_path)
    return True


def write_markdown(rows, charts, finds, out_path: pathlib.Path) -> None:
    lines = ["# Báo cáo kết quả tự động", "",
             "| " + " | ".join(COLS) + " |", "|" + "---|" * len(COLS)]
    for r in sorted(rows, key=lambda x: (x["dataset"], -x["average_accuracy"])):
        lines.append("| " + " | ".join(fmt(r.get(c), c) for c in COLS) + " |")
    lines += ["", "## Kết luận tự động"] + [f"- {f}" for f in finds]
    lines += ["", "## Biểu đồ"] + [f"![]({p})" for p in charts]
    out_path.write_text("\n".join(lines), encoding="utf-8")


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--dir", default="./artifacts")
    args = ap.parse_args()
    art = pathlib.Path(args.dir)
    rows = load_rows(art / "results")
    if not rows:
        print("Chưa có run nào trong artifacts/results — chạy run_all.py trước.")
        return 1
    charts = make_charts(rows, art / "report_charts")
    finds = auto_findings(rows)
    docx_path = art / "BAO_CAO_KET_QUA.docx"
    if write_docx(rows, charts, finds, docx_path):
        print(f"Đã ghi {docx_path}")
    else:
        md = art / "BAO_CAO_KET_QUA.md"
        write_markdown(rows, charts, finds, md)
        print(f"(Chưa cài python-docx -> đã ghi {md}. Cài: pip install python-docx rồi chạy lại.)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
